# streamlit_app.py
import streamlit as st
import pandas as pd
import json
import re
import time
from PyPDF2 import PdfReader
from docx import Document
from fpdf import FPDF
from openai import OpenAI

# -----------------------------
# 🔐 App Config & Secrets
# -----------------------------
st.set_page_config(page_title="GenAI Rights Dashboard", layout="wide")

# Password protection
APP_PASSWORD = st.secrets.get("app_password", "demo123")

if "authenticated" not in st.session_state:
    st.session_state.authenticated = False

if not st.session_state.authenticated:
    password_input = st.text_input("Enter app password:", type="password")
    if st.button("Login"):
        if password_input == APP_PASSWORD:
            st.session_state.authenticated = True
            st.success("✅ Access granted")
        else:
            st.error("❌ Incorrect password")
    st.stop()

# OpenAI client setup
try:
    OPENAI_API_KEY = st.secrets["openai"]["api_key"]
    if not OPENAI_API_KEY:
        st.error("❌ OpenAI API key is missing! Add it in Streamlit Secrets.")
        st.stop()
    client = OpenAI(api_key=OPENAI_API_KEY)
except Exception as e:
    st.error(f"Failed to initialize OpenAI client: {e}")
    st.stop()

MAX_CHARS = 8000

# -----------------------------
# Sample Contracts
# -----------------------------
sample_contracts = [
    {
        "filename": "EU_TVOD_Exclusive.pdf",
        "type": "pdf",
        "text": """Distributor grants Apple exclusive TVOD rights in the EU...
Music synchronization rights are excluded."""
    },
    {
        "filename": "US_SVOD_Holdback.docx",
        "type": "docx",
        "text": """Distributor grants Apple non-exclusive SVOD rights in the US...
No music rights included."""
    },
    {
        "filename": "Movie_A_Rights.docx",
        "type": "docx",
        "text": """Contract Title: Movie A Distribution Agreement
Rights Type: Streaming
Territory: United States
Exclusivity: Exclusive
License Start Date: 2026-01-01
License End Date: 2026-12-31
Holdbacks: None
Music Clearance: Cleared
Options: Renewal for 1 year"""
    }
]

# -----------------------------
# 1️⃣ Contract Selection
# -----------------------------
st.title("📺 GenAI Rights & Conflict Dashboard")
st.header("1️⃣ Upload or Select Contracts")

input_mode = st.radio("Choose input method:", ["Upload Contracts", "Use Sample Contracts"])
contracts_data = []

if input_mode == "Upload Contracts":
    uploaded_files = st.file_uploader("Upload PDF or DOCX contracts", type=["pdf","docx"], accept_multiple_files=True)
    if uploaded_files:
        for uploaded_file in uploaded_files:
            contract_text = ""
            if uploaded_file.type == "application/pdf":
                pdf = PdfReader(uploaded_file)
                for page in pdf.pages:
                    contract_text += (page.extract_text() or "") + "\n"
            elif uploaded_file.type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
                doc = Document(uploaded_file)
                for para in doc.paragraphs:
                    contract_text += para.text + "\n"
            contract_text = contract_text[:MAX_CHARS]
            contracts_data.append({"filename": uploaded_file.name, "text": contract_text})
else:
    all_files = [c["filename"] for c in sample_contracts]
    selected_files = st.multiselect("Select one or more sample contracts", options=all_files, default=all_files)
    for fname in selected_files:
        contract = next((c for c in sample_contracts if c["filename"] == fname), None)
        if contract:
            contracts_data.append({"filename": contract["filename"], "text": contract["text"]})

if not contracts_data:
    st.info("Upload or select contracts to begin analysis.")
    st.stop()

# -----------------------------
# 2️⃣ Structured Rights Extraction
# -----------------------------
st.header("2️⃣ Structured Rights Extraction")
rights_dfs = []

progress = st.progress(0)
total = len(contracts_data)

for i, contract in enumerate(contracts_data):
    st.info(f"Extracting rights from {contract['filename']}…")
    prompt = f"""
Extract rights attributes from this contract.

Return a SINGLE JSON object with EXACT keys:

{{
  "Rights Type": "",
  "Territory": "",
  "Exclusivity": "",
  "License Start Date": "",
  "License End Date": "",
  "Holdbacks": "",
  "Music Clearance": "",
  "Options": ""
}}

No markdown. No explanation.

Contract:
{contract['text']}
"""
    try:
        response = client.chat.completions.create(
            model="gpt-4o-mini",
            messages=[{"role":"system","content":"Return valid JSON only."},
                      {"role":"user","content":prompt}],
            response_format={"type":"json_object"}
        )
        raw_output = response.choices[0].message.content
        st.code(raw_output, language="json")
        try:
            parsed = json.loads(raw_output)
        except:
            match = re.search(r"\{.*\}", raw_output, re.DOTALL)
            parsed = json.loads(match.group()) if match else {}
        required_keys = ["Rights Type","Territory","Exclusivity",
                         "License Start Date","License End Date",
                         "Holdbacks","Music Clearance","Options"]
        for key in required_keys:
            if key not in parsed:
                parsed[key] = None
        df = pd.DataFrame([parsed])
        df["Contract"] = contract["filename"]
        rights_dfs.append(df)
    except Exception as e:
        st.error(f"Failed to extract rights for {contract['filename']}: {str(e)}")
    progress.progress((i+1)/total)
    time.sleep(0.5)

if not rights_dfs:
    st.warning("⚠️ No structured rights extracted.")
    st.stop()

combined_df = pd.concat(rights_dfs, ignore_index=True)
combined_df["License Start Date"] = pd.to_datetime(combined_df["License Start Date"], errors="coerce")
combined_df["License End Date"] = pd.to_datetime(combined_df["License End Date"], errors="coerce")

# -----------------------------
# 3️⃣ Conflict Detection
# -----------------------------
st.header("3️⃣ Conflict & Holdback Intelligence")
conflicts = []

for i in range(len(combined_df)):
    for j in range(i+1, len(combined_df)):
        r1 = combined_df.iloc[i]
        r2 = combined_df.iloc[j]
        territory_conflict = str(r1["Territory"]).lower() == str(r2["Territory"]).lower()
        date_conflict = pd.notna(r1["License Start Date"]) and pd.notna(r2["License Start Date"]) and \
                        r1["License Start Date"] <= r2["License End Date"] and \
                        r2["License Start Date"] <= r1["License End Date"]
        exclusive_conflict = str(r1["Exclusivity"]).lower()=="exclusive" or str(r2["Exclusivity"]).lower()=="exclusive"
        if territory_conflict and date_conflict and exclusive_conflict:
            conflicts.append(f"{r1['Contract']} ↔ {r2['Contract']} (Exclusive overlap in {r1['Territory']})")

# -----------------------------
# KPI Summary
# -----------------------------
st.subheader("📊 Summary")
col1, col2, col3, col4 = st.columns(4)
col1.metric("Total Contracts", len(combined_df))
col2.metric("Exclusive Deals", combined_df["Exclusivity"].astype(str).str.lower().eq("exclusive").sum())
col3.metric("Contracts with Holdbacks", combined_df["Holdbacks"].astype(str).str.lower().ne("none").sum())
col4.metric("Detected Conflicts", len(conflicts))

# -----------------------------
# Combined Rights Table
# -----------------------------
def highlight(row):
    holdback_flag = str(row["Holdbacks"]).strip().lower() not in ["none",""]
    conflict_flag = any(row["Contract"] in c for c in conflicts)
    colors = []
    for col in row.index:
        if conflict_flag:
            colors.append("background-color: #f28b82; color: black")
        elif holdback_flag:
            colors.append("background-color: #fff475; color: black")
        else:
            colors.append("background-color: #ccff90; color: black")
    return colors

st.subheader("Combined Rights Table")
st.dataframe(combined_df.style.apply(highlight, axis=1), width="stretch")

# -----------------------------
# Export CSV & PDF
# -----------------------------
csv_data = combined_df.to_csv(index=False).encode("utf-8")
st.download_button("📥 Download CSV", csv_data, "combined_rights.csv", "text/csv")

pdf = FPDF(orientation="L", unit="mm", format="A4")
pdf.set_auto_page_break(auto=True, margin=10)
pdf.add_page()
pdf.set_font("Arial", style="B", size=12)
pdf.cell(0, 10, "GenAI Rights Conflict Dashboard", ln=True)
pdf.ln(5)
pdf.set_font("Arial", size=8)

cols = combined_df.columns.tolist()
page_width = pdf.w - 20
col_width = page_width / len(cols)

# Header
pdf.set_font("Arial", style="B", size=8)
for c in cols:
    pdf.cell(col_width, 8, str(c), border=1)
pdf.ln()

pdf.set_font("Arial", size=8)
for _, row in combined_df.iterrows():
    for c in cols:
        pdf.multi_cell(col_width, 5, str(row[c]), border=1)
    pdf.ln()
pdf_bytes = pdf.output(dest="S").encode("latin1")
st.download_button("📥 Download PDF", pdf_bytes, "combined_rights.pdf", "application/pdf")

# -----------------------------
# 4️⃣ Auto-generated User Stories
# -----------------------------
st.header("4️⃣ Backlog User Stories")
for i, contract in enumerate(contracts_data):
    st.info(f"Generating user story for {contract['filename']}…")
    story_prompt = f"""
Create an Agile user story for a Product Owner.

Include ONLY:
- User Story
- Acceptance Criteria
- Test Notes

Contract:
{contract['text']}
"""
    try:
        story_response = client.chat.completions.create(
            model="gpt-4o-mini",
            messages=[{"role":"user","content":story_prompt}]
        )
        story = story_response.choices[0].message.content
        st.subheader(contract["filename"])
        st.text_area(f"User Story - {contract['filename']}", story, height=200)
    except Exception as e:
        st.error(f"Failed to generate story for {contract['filename']}: {e}")
    time.sleep(0.5)

st.success("🎉 Dashboard ready! CSV & PDF exports available.")